"""한성자동차 판매사원 파서/내보내기 테스트 (네트워크 불필요)."""

from __future__ import annotations

from pathlib import Path

from butimi import hansung_staff as hs

FIXTURE = Path(__file__).parent / "fixtures" / "hansung_store_sample.html"


def test_parse_staff_extracts_name_contact():
    html = FIXTURE.read_text(encoding="utf-8")
    members = hs.parse_staff(html, showroom="방배 전시장", detail_url="u")

    by_name = {m.name: m for m in members}
    assert set(by_name) == {"김철수", "이영희", "박민수"}

    # 소속 전시장이 모두 채워진다
    assert all(m.showroom == "방배 전시장" for m in members)

    # 첫 전화번호가 정규화되어 연락처로 들어간다(휴대폰 우선)
    assert by_name["김철수"].contact == "010-1234-5678"
    assert by_name["이영희"].contact == "010-2345-6789"  # 직통보다 앞선 휴대폰
    assert by_name["박민수"].contact == "010-3456-7890"  # 점 표기 → 하이픈


def test_stopwords_not_treated_as_names():
    # 직급/일반 단어는 이름으로 잡히지 않는다
    names = {m.name for m in hs.parse_staff(FIXTURE.read_text(encoding="utf-8"), "s")}
    assert "판매팀장" not in names
    assert "차장" not in names
    assert "메르세데스" not in names


def test_normalize_phone():
    assert hs.normalize_phone("01012345678") == "010-1234-5678"
    assert hs.normalize_phone("02-3479-8600") == "02-3479-8600"
    assert hs.normalize_phone("0234798600") == "02-3479-8600"
    assert hs.normalize_phone("031-710-8000") == "031-710-8000"


def test_discover_store_links():
    list_html = """
      <ul>
        <li><a href="/sales/retail-store-1">강남 전시장</a></li>
        <li><a href="/sales/retail-store-5">방배 전시장</a></li>
        <li><a href="/sales/retail-store-5">방배 전시장</a></li>
        <li><a href="/sales/other">무관 링크</a></li>
      </ul>
    """
    links = hs.discover_store_links(list_html)
    assert links == [
        ("강남 전시장", "https://mb.hansung.co.kr/sales/retail-store-1"),
        ("방배 전시장", "https://mb.hansung.co.kr/sales/retail-store-5"),
    ]


def test_to_xlsx_and_csv(tmp_path):
    members = [
        hs.StaffMember("김철수", "010-1234-5678", "방배 전시장"),
        hs.StaffMember("이영희", "010-2345-6789", "방배 전시장"),
    ]
    xlsx = hs.to_xlsx(members, tmp_path / "staff.xlsx")
    csv = hs.to_csv(members, tmp_path / "staff.csv")
    assert xlsx.exists() and csv.exists()

    from openpyxl import load_workbook
    ws = load_workbook(xlsx).active
    assert [c.value for c in ws[1]] == ["이름", "연락처", "전시장"]
    assert ws.max_row == 3  # 헤더 + 2명
